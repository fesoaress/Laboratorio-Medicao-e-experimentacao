"""Gera o DOCX acadêmico final a partir do relatório Markdown.

Execute da raiz do repositório:
    python -m lab02.reporting.build_final_report
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parents[2]
DEFAULT_MARKDOWN = BASE_DIR / "reports" / "final" / "lab02_relatorio_final.md"
DEFAULT_OUTPUT = BASE_DIR / "reports" / "final" / "lab02_relatorio_final.docx"

BLUE = "3333B2"
MID_BLUE = "4D4DC4"
LIGHT_BLUE = "D3D3EE"
VERY_LIGHT_BLUE = "F4F4FB"
GRAY = "5B6573"
LIGHT_GRAY = "D9DCE3"
WHITE = "FFFFFF"
BLACK = "202124"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, *, cover: bool = False) -> None:
    tbl_pr = table._tbl.tblPr
    current = tbl_pr.find(qn("w:tblBorders"))
    if current is not None:
        tbl_pr.remove(current)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        if cover or edge in {"left", "right", "insideV"}:
            node.set(qn("w:val"), "nil")
        else:
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), "8" if edge in {"top", "bottom"} else "4")
            node.set(qn("w:color"), BLUE if edge in {"top", "bottom"} else LIGHT_GRAY)
        borders.append(node)
    tbl_pr.append(borders)


def style_section_heading(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), BLUE)
    p_pr.append(shading)
    paragraph.paragraph_format.left_indent = Cm(0.18)
    paragraph.paragraph_format.right_indent = Cm(0)
    paragraph.paragraph_format.space_before = Pt(13)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.08
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor.from_string(WHITE)


def style_result_box(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), VERY_LIGHT_BLUE)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "22")
    left.set(qn("w:space"), "10")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.25)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_together = True


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, value, end):
        run._r.append(element)


def set_repeatable_header_footer(document: Document) -> None:
    section = document.sections[0]
    section.different_first_page_header_footer = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.text = "PUC Minas  •  Medição e Experimentação de Software  •  Laboratório 02"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.style = document.styles["Header"]

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("PUC Minas  |  Página ")
    run.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(paragraph)


def set_document_defaults(document: Document) -> None:
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size, color in (
        ("Title", 22, BLUE),
        ("Heading 1", 14, WHITE),
        ("Heading 2", 12, BLUE),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    styles["Heading 1"].paragraph_format.keep_together = True

    for style_name in ("Header", "Footer"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(8)
        style.font.color.rgb = RGBColor.from_string(GRAY)

    if "Figure Caption" not in styles:
        caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Figure Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_together = True

    if "Table Caption" not in styles:
        table_caption = styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_caption = styles["Table Caption"]
    table_caption.font.name = "Arial"
    table_caption.font.size = Pt(9)
    table_caption.font.bold = True
    table_caption.font.color.rgb = RGBColor.from_string(BLUE)
    table_caption.paragraph_format.space_before = Pt(7)
    table_caption.paragraph_format.space_after = Pt(3)
    table_caption.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8.5)
    code_style.paragraph_format.left_indent = Cm(0.5)
    code_style.paragraph_format.right_indent = Cm(0.5)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(6)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


INLINE_RE = re.compile(
    r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\)|\*[^*]+?\*)"
)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MID_BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline_markdown(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(12)

    band = document.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    band.autofit = False
    set_table_borders(band, cover=True)
    cell = band.cell(0, 0)
    cell.width = Cm(16.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=420, start=220, bottom=420, end=220)
    set_cell_shading(cell, BLUE)

    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("LABORATÓRIO 02")
    run.font.name = "Arial"
    run.font.size = Pt(23)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(WHITE)

    paragraph = cell.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Assistentes de IA vs. Codificação Manual")
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(WHITE)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(28)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead = paragraph.add_run("Participantes:")
    lead.font.bold = True
    lead.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph.paragraph_format.space_after = Pt(5)
    for participant in ("Fernanda Soares", "Islayder Jackson", "Vinicius Gomes"):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(participant)
        paragraph.paragraph_format.space_after = Pt(2)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(20)

    for text, size, bold, color in (
        ("Pontifícia Universidade Católica de Minas Gerais", 11, True, BLUE),
        ("Engenharia de Software", 10.5, True, BLACK),
        ("Medição e Experimentação de Software", 10.5, False, GRAY),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        paragraph.paragraph_format.space_after = Pt(3)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead = paragraph.add_run("Professor: ")
    lead.font.bold = True
    paragraph.add_run("Danilo Maia")
    lead.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph.paragraph_format.space_before = Pt(16)
    paragraph.paragraph_format.space_after = Pt(4)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(38)
    run = paragraph.add_run("Setembro de 2026")
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.bold = True

    document.add_page_break()


def clean_cell(text: str) -> str:
    return text.strip().replace("**", "").replace("`", "")


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        row = [clean_cell(cell) for cell in lines[index].strip().strip("|").split("|")]
        rows.append(row)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    font_size = 7.0 if columns >= 8 else 8.0 if columns >= 6 else 8.7

    for row_index, source_row in enumerate(rows):
        target_row = table.rows[row_index]
        set_row_cant_split(target_row)
        if row_index == 0:
            set_repeat_table_header(target_row)
        for column_index, cell in enumerate(target_row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = source_row[column_index] if column_index < len(source_row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = paragraph.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(font_size)
            if row_index == 0:
                set_cell_shading(cell, BLUE)
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string(WHITE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, VERY_LIGHT_BLUE)
            if row_index == len(rows) - 1 and text.lower() == "total":
                run.font.bold = True
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_result_box(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    style_result_box(paragraph)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        add_inline_markdown(paragraph, line)
    for run in paragraph.runs:
        run.font.name = "Arial"
        if run.bold:
            run.font.color.rgb = RGBColor.from_string(BLUE)


def add_picture(document: Document, path: Path, alt_text: str) -> None:
    if not path.exists():
        raise FileNotFoundError(f"Figura não encontrada: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Cm(16.0))
    shape._inline.docPr.set("descr", alt_text)


def add_code_block(document: Document, code_lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run("\n".join(code_lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F2F4F7")
    p_pr.append(shading)


def render_markdown(document: Document, markdown_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    start = next(index for index, line in enumerate(lines) if line.startswith("## 2. "))
    index = start
    paragraph_buffer: list[str] = []
    in_code = False
    code_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        text = " ".join(item.strip() for item in paragraph_buffer).strip()
        paragraph_buffer.clear()
        if not text:
            return
        style = None
        if text.startswith("**Tabela ") and text.endswith("**"):
            style = "Table Caption"
            text = text[2:-2]
        elif text.startswith("*Figura ") and text.endswith("*"):
            style = "Figure Caption"
            text = text[1:-1]
        paragraph = document.add_paragraph(style=style)
        add_inline_markdown(paragraph, text)

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            heading = stripped[3:]
            heading_paragraph = document.add_heading(heading, level=1)
            if heading.startswith(("13.", "15.", "18.")):
                heading_paragraph.paragraph_format.page_break_before = True
            style_section_heading(heading_paragraph)
            index += 1
            continue
        if stripped.startswith(">"):
            flush_paragraph()
            box_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                content = lines[index].strip()[1:].strip()
                if content:
                    box_lines.append(content)
                index += 1
            add_result_box(document, box_lines)
            continue
        if stripped.startswith("!["):
            flush_paragraph()
            match = re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)", stripped)
            if not match:
                raise ValueError(f"Sintaxe de figura inválida: {stripped}")
            alt_text, relative_path = match.groups()
            add_picture(document, (markdown_path.parent / relative_path).resolve(), alt_text)
            index += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Number")
            add_inline_markdown(paragraph, ordered.group(2))
            index += 1
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, stripped[2:])
            index += 1
            continue
        paragraph_buffer.append(stripped.rstrip("  "))
        index += 1

    flush_paragraph()
    if in_code:
        add_code_block(document, code_lines)


def build(markdown_path: Path, output_path: Path) -> tuple[int, int, int]:
    document = Document()
    set_document_defaults(document)
    set_repeatable_header_footer(document)

    properties = document.core_properties
    properties.title = "Laboratório 02 — Assistentes de IA vs. Codificação Manual"
    properties.subject = "Relatório final de Medição e Experimentação de Software"
    properties.author = "Fernanda Soares; Islayder Jackson; Vinicius Gomes"
    properties.keywords = "Lab02, IA, codificação manual, medição, experimentação"

    add_cover(document)
    render_markdown(document, markdown_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)

    validation = Document(output_path)
    headings = sum(
        1 for paragraph in validation.paragraphs if paragraph.style.name.startswith("Heading")
    )
    tables = len(validation.tables) - 1  # exclui o bloco visual da capa
    figures = len(validation.inline_shapes)
    if headings != 20:
        raise RuntimeError(f"Esperados 20 títulos após a capa; encontrados {headings}")
    if tables != 10:
        raise RuntimeError(f"Esperadas 10 tabelas; encontradas {tables}")
    if figures != 9:
        raise RuntimeError(f"Esperadas 9 figuras; encontradas {figures}")
    return headings, tables, figures


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--markdown", type=Path, default=DEFAULT_MARKDOWN)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    headings, tables, figures = build(args.markdown.resolve(), args.output.resolve())
    print(f"DOCX gerado: {args.output.resolve()}")
    print(f"Validação estrutural: {headings} seções, {tables} tabelas, {figures} figuras")


if __name__ == "__main__":
    main()
