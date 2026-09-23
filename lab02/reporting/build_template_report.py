"""Gera o relatório final do Lab02 a partir do template da disciplina."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE = Path.home() / "Downloads" / "Template_Relatorio_Laboratorio.docx"
DEFAULT_MARKDOWN = BASE_DIR / "reports" / "final" / "lab02_relatorio_final.md"
DEFAULT_OUTPUT = BASE_DIR / "reports" / "final" / "lab02_relatorio_final.docx"

NAVY = "1F3A5F"
GREEN = "1F6E63"
LIGHT_GREEN = "EAF3F1"
LIGHT_GRAY = "F2F4F6"
WHITE = "FFFFFF"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


def clear_body(document: Document) -> None:
    """Remove o conteúdo demonstrativo, preservando estilos e configuração."""
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_row_cant_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    for kind, text in (
        ("begin", None),
        (None, " PAGE "),
        ("separate", None),
        (None, "1"),
        ("end", None),
    ):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        elif text == " PAGE ":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        else:
            element = OxmlElement("w:t")
            element.text = text
        run._r.append(element)


def configure_document(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.paragraph_format.keep_with_next = True

    section = document.sections[0]
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    header = section.header.paragraphs[0]
    header.text = "PUC Minas  •  Laboratório de Experimentação de Software  •  Lab02"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(NAVY)

    footer = section.footer.paragraphs[0]
    footer.text = "Laboratório 02  |  Página "
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(NAVY)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_front_matter(document: Document) -> None:
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Relatório de Laboratório")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("Laboratório 02 — Assistentes de IA vs. Codificação Manual")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(GREEN)

    rows = (
        ("Curso", "Engenharia de Software"),
        ("Disciplina", "Laboratório de Experimentação de Software"),
        ("Turno / Período", "Noite / 6º"),
        ("Professor", "Danilo Maia"),
        ("Laboratório", "Lab02 — Assistentes de IA vs. Codificação Manual"),
        ("Grupo", "Fernanda Soares · Islayder Jackson · Vinicius Gomes"),
        (
            "Repositório",
            "https://github.com/fesoaress/Laboratorio-Medicao-e-experimentacao",
        ),
        ("Data de entrega", "23 de setembro de 2026"),
    )
    table = document.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, (label, value) in enumerate(rows):
        row = table.rows[index]
        set_row_cant_split(row)
        left, right = row.cells
        left.text = label
        right.text = value
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(left, NAVY)
        for run in left.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
        if index % 2:
            set_cell_shading(right, LIGHT_GRAY)
        for cell in row.cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)

    document.add_page_break()


def add_inline_markdown(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(GREEN)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def clean_cell(text: str) -> str:
    return text.strip().replace("**", "").replace("`", "")


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        row = [clean_cell(cell) for cell in lines[index].strip().strip("|").split("|")]
        rows.append(row)
        index += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    font_size = 6.6 if column_count >= 8 else 7.4 if column_count >= 6 else 8.2

    for row_index, source in enumerate(rows):
        row = table.rows[row_index]
        set_row_cant_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            value = source[column_index] if column_index < len(source) else ""
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if column_index == 0
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            run = paragraph.runs[0]
            run.font.size = Pt(font_size)
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(WHITE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT_GREEN)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_picture(document: Document, path: Path, alt_text: str) -> None:
    if not path.exists():
        raise FileNotFoundError(f"Figura não encontrada: {path}")
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Cm(15.8))
    shape._inline.docPr.set("descr", alt_text)


def render_markdown(document: Document, markdown_path: Path) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = next(i for i, line in enumerate(lines) if line.startswith("## 1. "))
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_buffer:
            return
        text = " ".join(item.strip() for item in paragraph_buffer).strip()
        paragraph_buffer.clear()
        if not text:
            return
        style = "Caption" if text.startswith("*Figura ") and text.endswith("*") else None
        if style:
            text = text[1:-1]
        paragraph = document.add_paragraph(style=style)
        add_inline_markdown(paragraph, text)
        if style:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    while index < len(lines):
        stripped = lines[index].strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1)) - 1
            text = heading.group(2)
            paragraph = document.add_heading(text, level=level)
            if level == 1 and text.startswith(("4.", "5.", "6.")):
                paragraph.paragraph_format.page_break_before = True
            index += 1
            continue
        if stripped.startswith("!["):
            flush_paragraph()
            match = re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)", stripped)
            if not match:
                raise ValueError(f"Sintaxe de figura inválida: {stripped}")
            alt_text, relative_path = match.groups()
            add_picture(document, (markdown_path.parent / relative_path).resolve(), alt_text)
            index += 1
            continue
        if stripped.startswith("|"):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, stripped[2:])
            index += 1
            continue
        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()


def validate(document_path: Path) -> tuple[int, int, int]:
    document = Document(document_path)
    headings = sum(
        1 for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")
    )
    figures = len(document.inline_shapes)
    tables = len(document.tables)
    complete_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    complete_text += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )

    forbidden = (
        "youtube.com",
        "delegated",
        "não fez",
        "fez errado",
        "não realizou",
    )
    found = [term for term in forbidden if term.casefold() in complete_text.casefold()]
    if found:
        raise RuntimeError(f"Conteúdo proibido encontrado: {', '.join(found)}")
    if headings < 12:
        raise RuntimeError(f"Estrutura incompleta: somente {headings} títulos")
    if tables < 8:
        raise RuntimeError(f"Estrutura incompleta: somente {tables} tabelas")
    if figures != 8:
        raise RuntimeError(f"Esperadas 8 figuras; encontradas {figures}")
    return headings, tables, figures


def build(template_path: Path, markdown_path: Path, output_path: Path) -> tuple[int, int, int]:
    if not template_path.exists():
        raise FileNotFoundError(f"Template não encontrado: {template_path}")
    document = Document(template_path)
    clear_body(document)
    configure_document(document)

    properties = document.core_properties
    properties.title = "Laboratório 02 — Assistentes de IA vs. Codificação Manual"
    properties.subject = "Relatório final de Laboratório de Experimentação de Software"
    properties.author = "Fernanda Soares; Islayder Jackson; Vinicius Gomes"
    properties.keywords = "Lab02, inteligência artificial, codificação manual, experimento"

    add_front_matter(document)
    render_markdown(document, markdown_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return validate(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--markdown", type=Path, default=DEFAULT_MARKDOWN)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    headings, tables, figures = build(
        args.template.resolve(),
        args.markdown.resolve(),
        args.output.resolve(),
    )
    print(f"DOCX gerado: {args.output.resolve()}")
    print(f"Validação estrutural: {headings} títulos, {tables} tabelas, {figures} figuras")


if __name__ == "__main__":
    main()
